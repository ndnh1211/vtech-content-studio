"""Bước 1: DOCX -> content.html + seo.json + images/ (offline, chỉ đọc DOCX).
Dùng: python tools/docx_extract.py --docx "<file.docx>" --out content/<slug>/source
"""
import argparse, json, re, sys
from pathlib import Path
import docx
from docx.text.run import Run
from docx.text.hyperlink import Hyperlink
from html import escape

SEO_LABELS = {
    "focus keyphrase": "focus_keyphrase", "cụm từ khóa chính": "focus_keyphrase",
    "tiêu đề seo": "seo_title", "đường dẫn (slug)": "slug", "slug": "slug",
    "mô tả meta": "meta_description", "alt text ảnh": "image_alt",
}


def inline_html(p):
    out = []
    for item in p.iter_inner_content():
        if isinstance(item, Hyperlink):
            txt = escape(item.text)
            out.append(f'<a href="{escape(item.url or "", quote=True)}">{txt}</a>' if item.url else txt)
        elif isinstance(item, Run) and item.text:
            t = escape(item.text)
            out.append(f"<strong>{t}</strong>" if item.bold else t)
    return "".join(out).strip()


def is_list(p):
    return "list" in (p.style.name or "").lower() or p._p.pPr is not None and p._p.pPr.numPr is not None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    d = docx.Document(a.docx)
    out = Path(a.out)
    (out / "images").mkdir(parents=True, exist_ok=True)

    seo = {}
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            key = SEO_LABELS.get(cells[0].lower()) if len(cells) > 1 else None
            if key:
                seo[key] = cells[1]

    html, h1, images, in_ul = [], "", [], False
    first_para = ""
    for p in d.paragraphs:
        # ảnh nằm trong đoạn
        for rid in p._p.xpath(".//a:blip/@r:embed"):
            part = d.part.related_parts[rid]
            name = f"image{len(images)+1}{Path(part.partname).suffix}"
            (out / "images" / name).write_bytes(part.blob)
            images.append(name)
            if in_ul: html.append("</ul>"); in_ul = False
            html.append(f"<!--image:{name}-->")
        text = inline_html(p)
        if not text: continue
        style = (p.style.name or "") if p.style else ""
        if style == "Heading 1" and not h1:
            h1 = p.text.strip(); continue
        if is_list(p):
            if not in_ul: html.append("<ul>"); in_ul = True
            html.append(f"<li>{text}</li>"); continue
        if in_ul: html.append("</ul>"); in_ul = False
        m = re.match(r"Heading (\d)", style)
        if m: html.append(f"<h{m[1]}>{text}</h{m[1]}>")
        elif p.text.strip().upper().startswith("THÔNG TIN SEO"): continue
        else:
            html.append(f"<p>{text}</p>")
            first_para = first_para or p.text.strip()
    if in_ul: html.append("</ul>")

    (out / "content.html").write_text("\n".join(html), encoding="utf-8")
    meta = {"h1": h1, "seo": seo, "images": images, "first_paragraph": first_para,
            "source_docx": a.docx}
    (out / "seo.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK: h1={bool(h1)} seo_fields={len(seo)} images={len(images)} blocks={len(html)} -> {out}")
    missing = {"focus_keyphrase", "seo_title", "slug", "meta_description"} - set(seo)
    if missing: print("CẢNH BÁO thiếu SEO:", ", ".join(sorted(missing)))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
